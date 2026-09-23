from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient

from app.main import app


def test_meeting_mock_flow():
    with TestClient(app) as client:
        home = client.get("/")
        health = client.get("/health")
        assert home.status_code == 200 and "Aqylman" in home.text
        assert health.status_code == 200 and health.json()["mode"] == "demo"

        created = client.post("/api/v1/meetings", json={"title": "Тест", "meeting_date": "2026-09-23"})
        assert created.status_code == 201
        meeting_id = created.json()["id"]

        uploaded = client.post(
            f"/api/v1/meetings/{meeting_id}/upload",
            files={"file": ("meeting.wav", b"mock audio", "audio/wav")},
        )
        assert uploaded.status_code == 200

        processed = client.post(f"/api/v1/meetings/{meeting_id}/process")
        assert processed.status_code == 202
        result = client.get(f"/api/v1/meetings/{meeting_id}").json()
        assert result["status"] == "completed"
        assert len(result["tasks"]) == 2
        assert result["summary"]["topic"]

        renamed = client.patch(
            f"/api/v1/meetings/{meeting_id}/participants/{result['participants'][0]['id']}",
            json={"display_name": "Әлия Қанатқызы", "role": "Өндіріс қауіпсіздігі бөлімінің жетекшісі"},
        )
        assert renamed.status_code == 200
        assert renamed.json()["display_name"] == "Әлия Қанатқызы"
        assert renamed.json()["role"] == "Өндіріс қауіпсіздігі бөлімінің жетекшісі"
        updated_result = client.get(f"/api/v1/meetings/{meeting_id}").json()
        assert updated_result["tasks"][0]["assigned_by"] == "Әлия Қанатқызы"

        pdf = client.get(f"/api/v1/meetings/{meeting_id}/export/pdf")
        docx = client.get(f"/api/v1/meetings/{meeting_id}/export/docx")
        assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
        assert docx.status_code == 200 and docx.content.startswith(b"PK")
        document = Document(BytesIO(docx.content))
        assert len(document.tables) >= 2
        full_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        assert "ПОЛНЫЙ ТРАНСКРИПТ" in full_text
        assert "Әлия Қанатқызы" in full_text
