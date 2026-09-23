from __future__ import annotations

import io
import json
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.models.meeting import Meeting


def stamp(seconds: float) -> str:
    value = max(0, int(seconds))
    return f"{value // 3600:02}:{value // 60 % 60:02}:{value % 60:02}"


class ExportService:
    @staticmethod
    def _json(value: str) -> list[str]:
        try:
            return json.loads(value)
        except (TypeError, ValueError):
            return []

    def docx(self, meeting: Meeting) -> bytes:
        document = Document()
        normal = document.styles["Normal"]
        normal.font.name = "Arial"
        normal.font.size = Pt(10)
        title = document.add_heading("ПРОТОКОЛ СОВЕЩАНИЯ", 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        document.add_paragraph(f"Название: {meeting.title}")
        document.add_paragraph(f"Дата: {meeting.meeting_date.isoformat()}")
        document.add_paragraph(f"Тема: {meeting.summary.topic if meeting.summary else 'Не определена'}")

        document.add_heading("УЧАСТНИКИ", 1)
        people = document.add_table(rows=1, cols=3)
        people.style = "Table Grid"
        for cell, value in zip(people.rows[0].cells, ["№", "Имя", "Должность / роль"]):
            cell.text = value
        for index, person in enumerate(meeting.participants, 1):
            cells = people.add_row().cells
            cells[0].text, cells[1].text, cells[2].text = str(index), person.display_name, person.role or "Не указана"

        summary = meeting.summary
        document.add_heading("КРАТКОЕ СОДЕРЖАНИЕ", 1)
        document.add_paragraph(summary.summary_text if summary and summary.summary_text else "Не сформировано")
        for heading, values in (
            ("КЛЮЧЕВЫЕ ВОПРОСЫ", self._json(summary.key_points_json) if summary else []),
            ("ПРОБЛЕМЫ", self._json(summary.problems_json) if summary else []),
            ("ПРИНЯТЫЕ РЕШЕНИЯ", self._json(summary.decisions_json) if summary else []),
        ):
            document.add_heading(heading, 1)
            if values:
                for value in values:
                    document.add_paragraph(value, style="List Bullet")
            else:
                document.add_paragraph("Не указаны")

        document.add_heading("ПОРУЧЕНИЯ", 1)
        tasks = document.add_table(rows=1, cols=6)
        tasks.style = "Table Grid"
        headers = ["№", "Поручение", "Ответственный", "Кто поставил", "Срок", "Статус"]
        for cell, value in zip(tasks.rows[0].cells, headers):
            cell.text = value
        states = {"open": "В работе", "confirmed": "Проверено", "done": "Выполнено"}
        for index, task in enumerate(meeting.tasks, 1):
            values = [str(index), task.task, task.responsible, task.assigned_by,
                      str(task.deadline_normalized or task.deadline_raw or "Не указан"), states.get(task.status, task.status)]
            for cell, value in zip(tasks.add_row().cells, values):
                cell.text = value

        document.add_page_break()
        document.add_heading("ПОЛНЫЙ ТРАНСКРИПТ", 1)
        for segment in sorted(meeting.transcript, key=lambda item: item.start):
            document.add_paragraph(f"[{stamp(segment.start)} – {stamp(segment.end)}]")
            name = document.add_paragraph()
            name.add_run(segment.speaker_name).bold = True
            if segment.speaker_role:
                document.add_paragraph(segment.speaker_role)
            document.add_paragraph(segment.text)
            document.add_paragraph("—" * 28)
        output = io.BytesIO()
        document.save(output)
        return output.getvalue()

    @staticmethod
    def _font() -> str:
        candidates = [
            "C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/DejaVuSans.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            "/System/Library/Fonts/Supplemental/Arial.ttf",
        ]
        path = next((Path(item) for item in candidates if Path(item).is_file()), None)
        if not path:
            raise RuntimeError("EXPORT_FAILED: Не найден локальный Unicode-шрифт для PDF.")
        if "ProtocolUnicode" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("ProtocolUnicode", str(path)))
        return "ProtocolUnicode"

    def pdf(self, meeting: Meeting) -> bytes:
        output = io.BytesIO()
        font = self._font()
        styles = getSampleStyleSheet()
        for style in styles.byName.values():
            style.fontName = font
        story = [Paragraph("ПРОТОКОЛ СОВЕЩАНИЯ", styles["Title"]), Spacer(1, 4 * mm)]
        summary = meeting.summary
        for value in (f"Название: {meeting.title}", f"Дата: {meeting.meeting_date.isoformat()}",
                      f"Тема: {summary.topic if summary else 'Не определена'}"):
            story.append(Paragraph(escape(value), styles["BodyText"]))
        people_table = Table(
            [["№", "Имя", "Должность / роль"]] + [[str(i), p.display_name, p.role or "Не указана"] for i, p in enumerate(meeting.participants, 1)],
            colWidths=[12 * mm, 55 * mm, 100 * mm], repeatRows=1,
        )
        people_table.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), font), ("GRID", (0, 0), (-1, -1), .4, colors.grey)]))
        story += [Paragraph("УЧАСТНИКИ", styles["Heading2"]), people_table]
        story.append(Paragraph("КРАТКОЕ СОДЕРЖАНИЕ", styles["Heading2"]))
        story.append(Paragraph(escape(summary.summary_text if summary and summary.summary_text else "Не сформировано"), styles["BodyText"]))
        for heading, values in (("КЛЮЧЕВЫЕ ВОПРОСЫ", self._json(summary.key_points_json) if summary else []),
                                ("ПРОБЛЕМЫ", self._json(summary.problems_json) if summary else []),
                                ("ПРИНЯТЫЕ РЕШЕНИЯ", self._json(summary.decisions_json) if summary else [])):
            story.append(Paragraph(heading, styles["Heading2"]))
            for value in values or ["Не указаны"]:
                story.append(Paragraph("• " + escape(value), styles["BodyText"]))
        task_data = [["№", "Поручение", "Ответственный", "Срок"]] + [
            [str(i), Paragraph(escape(t.task), styles["BodyText"]), t.responsible,
             str(t.deadline_normalized or t.deadline_raw or "—")] for i, t in enumerate(meeting.tasks, 1)
        ]
        task_table = Table(task_data, colWidths=[10 * mm, 80 * mm, 48 * mm, 30 * mm], repeatRows=1)
        task_table.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), font), ("GRID", (0, 0), (-1, -1), .4, colors.grey),
                                       ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8efe9")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        story += [Paragraph("ПОРУЧЕНИЯ", styles["Heading2"]), task_table, PageBreak(), Paragraph("ПОЛНЫЙ ТРАНСКРИПТ", styles["Heading2"])]
        for segment in sorted(meeting.transcript, key=lambda item: item.start):
            story.append(Paragraph(f"[{stamp(segment.start)} – {stamp(segment.end)}]", styles["BodyText"]))
            story.append(Paragraph(escape(segment.speaker_name), styles["Heading3"]))
            if segment.speaker_role:
                story.append(Paragraph(escape(segment.speaker_role), styles["Italic"]))
            story.append(Paragraph(escape(segment.text), styles["BodyText"]))
            story.append(Spacer(1, 4 * mm))
        SimpleDocTemplate(output, pagesize=A4, rightMargin=15 * mm, leftMargin=15 * mm).build(story)
        return output.getvalue()
